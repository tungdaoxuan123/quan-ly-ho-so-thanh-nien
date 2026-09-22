import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import quan_ly_ho_so as app
from tao_ho_so_word import generate_document


TEMPLATE = Path(__file__).resolve().parent / "Mau_Ho_So_Thanh_Nien.docx"


def create_test_workbook(path):
    headers = ["STT"]
    for _, section_headers in app.FORM_GROUPS:
        for header in section_headers:
            if header not in headers:
                headers.append(header)
    records = [
        {
            "STT": 1,
            "TÊN THƯỜNG DÙNG": "NGUYỄN VĂN AN",
            "Tên khai sinh": "Nguyễn Văn An",
            "Năm sinh": 2005,
            "Căn cước": "001205000001",
            "Nghề nghiệp": "Sinh viên",
            "Văn hóa": "12/12",
            "Dân tộc": "Kinh",
            "Tôn giáo": "Không",
            "Thường trú": "Hà Nội",
        },
        {
            "STT": 2,
            "TÊN THƯỜNG DÙNG": "TRẦN THỊ BÌNH",
            "Năm sinh": 2004,
            "Căn cước": "001204000002",
            "Nghề nghiệp": "Nông dân",
            "Văn hóa": "12/12",
            "Dân tộc": "kinh",
            "Tôn giáo": "không",
            "Nơi ở hiện nay": "Đà Nẵng",
        },
        {
            "STT": 3,
            "TÊN THƯỜNG DÙNG": "PHAN TRẦN GIA HUY",
            "Năm sinh": 2003,
            "Căn cước": "001203000003",
            "Nghề nghiệp": "SINH VIÊN",
            "Văn hóa": "Đại học",
            "Dân tộc": "KINH",
            "Tôn giáo": "Không",
            "Thường trú": "TP Hồ Chí Minh",
        },
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hồ sơ"
    sheet.append(headers)
    sheet.append([])
    for record in records:
        sheet.append([record.get(header) for header in headers])
    workbook.save(path)
    workbook.close()
    return len(headers)


class QuanLyHoSoTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "records.xlsx"
        self.previous_database = app.state["database"]
        app.state["database"] = Path(self.temp_dir.name) / "cache.sqlite3"
        self.header_count = create_test_workbook(self.workbook_path)
        app.set_workbook(self.workbook_path, persist=False)

    def tearDown(self):
        app.state.update(
            workbook=None,
            sheet=None,
            headers={},
            database=self.previous_database,
            record_count=0,
            signature=None,
            loaded_at=None,
            error=None,
            last_output_dir=None,
        )
        self.temp_dir.cleanup()

    def test_reads_records_and_accent_insensitive_search(self):
        self.assertEqual(app.state["record_count"], 3)
        self.assertNotIn("records", app.state)
        matches, count = app.query_records("nguyen", {}, limit=50, offset=0)
        self.assertEqual(count, 1)
        self.assertEqual([record["stt"] for record in matches], ["1"])

        accented_matches, accented_count = app.query_records("Nguyễn", {}, limit=50, offset=0)
        self.assertEqual(accented_count, 1)
        self.assertEqual([record["stt"] for record in accented_matches], ["1"])

    def test_quick_search_does_not_match_unrelated_profile_fields(self):
        matches, count = app.query_records("Ha Noi", {}, limit=50, offset=0)
        self.assertEqual(count, 0)
        self.assertEqual(matches, [])

    def test_filters_ignore_case_and_accents(self):
        occupation = app.query_filter_options("occupation")[0]
        exact_matches, exact_count = app.query_records("", {"occupation": occupation}, limit=500, offset=0)
        normalized_matches, normalized_count = app.query_records(
            "", {"occupation": app.fold(occupation)}, limit=500, offset=0
        )
        self.assertGreater(exact_count, 0)
        self.assertEqual(exact_count, normalized_count)
        self.assertEqual([record["stt"] for record in exact_matches], [record["stt"] for record in normalized_matches])

    def test_filter_options_group_case_variations(self):
        options = app.query_filter_options("ethnicity")
        kinh_options = [value for value in options if app.fold(value) == "kinh"]
        self.assertEqual(len(kinh_options), 1)

    def test_external_excel_change_is_synchronized(self):
        workbook = load_workbook(self.workbook_path)
        name_column = app.state["headers"]["TÊN THƯỜNG DÙNG"]
        workbook.active.cell(3, name_column).value = "UPDATED OUTSIDE APP"
        workbook.save(self.workbook_path)
        workbook.close()

        self.assertTrue(app.refresh_state())
        self.assertEqual(app.find_record("1")["name"], "UPDATED OUTSIDE APP")

    def test_failed_sync_keeps_previous_sqlite_data(self):
        self.workbook_path.write_bytes(b"not an Excel workbook")
        self.assertFalse(app.refresh_state(force=True))
        self.assertEqual(app.database_record_count(), 3)
        self.assertIsNotNone(app.find_record("1"))

    def test_edit_and_add_preserve_workbook_shape(self):
        signature = app.signature_token(app.state["signature"])
        name_column = app.state["headers"]["TÊN THƯỜNG DÙNG"]
        row, stt, backup = app.save_person(
            self.workbook_path, app.state["sheet"], {name_column: "EDITED TEST"}, signature, original_stt="1"
        )
        self.assertEqual((row, stt), (3, "1"))
        self.assertTrue(backup.is_file())

        app.set_workbook(self.workbook_path, persist=False)
        signature = app.signature_token(app.state["signature"])
        row, stt, backup = app.save_person(
            self.workbook_path, app.state["sheet"], {name_column: "ADDED TEST"}, signature
        )
        self.assertEqual((row, stt), (6, "4"))
        self.assertTrue(backup.is_file())
        workbook = load_workbook(self.workbook_path, read_only=True, data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet.cell(3, name_column).value, "EDITED TEST")
        self.assertEqual(sheet.cell(6, name_column).value, "ADDED TEST")
        self.assertEqual(sheet.cell(6, app.state["headers"]["STT"]).value, 4)
        self.assertEqual(sheet.max_column, self.header_count)
        workbook.close()

    def test_stale_form_is_rejected(self):
        signature = app.signature_token(app.state["signature"])
        workbook = load_workbook(self.workbook_path)
        workbook.active["C3"] = "CHANGED OUTSIDE APP"
        workbook.save(self.workbook_path)
        workbook.close()
        with self.assertRaises(app.StaleWorkbookError):
            app.save_person(
                self.workbook_path,
                app.state["sheet"],
                {app.state["headers"]["TÊN THƯỜNG DÙNG"]: "STALE EDIT"},
                signature,
                original_stt="1",
            )

    def test_browser_routes_and_word_generation(self):
        client = app.app.test_client()
        with client:
            response = client.get("/")
            self.assertEqual(response.status_code, 200)
            self.assertIn(b"PHAN TR", response.data)
            page = response.get_data(as_text=True)
            self.assertIn("Tìm kiếm hồ sơ", page)
            self.assertIn("Bộ lọc nâng cao", page)
            self.assertIn("Thêm hồ sơ", page)
            response = client.get("/?q=nguyen&birth_year=2005")
            self.assertEqual(response.status_code, 200)
            self.assertIn('<details class="advanced" open>', response.get_data(as_text=True))
            form = client.get("/person/1/edit")
            self.assertEqual(form.status_code, 200)
            form_text = form.get_data(as_text=True)
            self.assertIn("Cập nhật hồ sơ", form_text)
            self.assertIn("Thông tin cá nhân", form_text)
            self.assertIn("Lưu vào Excel", form_text)
            self.assertIn('<details class="section-card" open>', form_text)
            csrf = re.search(r'name="csrf_token" value="([^"]+)"', form_text).group(1)
            response = client.post("/generate/1", data={"csrf_token": csrf}, follow_redirects=False)
            self.assertEqual(response.status_code, 302)
        output_dir = self.workbook_path.parent / app.DEFAULT_OUTPUT_NAME
        first = generate_document(self.workbook_path, app.state["sheet"], 3, TEMPLATE, output_dir)
        second = generate_document(self.workbook_path, app.state["sheet"], 3, TEMPLATE, output_dir)
        self.assertTrue(first.is_file())
        self.assertTrue(second.is_file())
        self.assertNotEqual(first, second)


if __name__ == "__main__":
    unittest.main()
