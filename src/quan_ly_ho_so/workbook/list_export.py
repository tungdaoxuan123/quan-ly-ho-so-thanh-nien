"""Fill the standard youth-list Excel template with the records picked in the table."""

import copy

from openpyxl import load_workbook

from quan_ly_ho_so.config import NUMERIC_HEADERS
from quan_ly_ho_so.errors import WorkbookError
from quan_ly_ho_so.utils.text import normalized_header
from quan_ly_ho_so.workbook.writer import copy_row_style

DIEN_HEADER = "Diện"
QUARTER_HEADER = "Khu phố"
# Columns the app owns rather than the workbook: (heading, record key), appended in this order.
APP_COLUMNS = ((QUARTER_HEADER, "quarter"), (DIEN_HEADER, "dien"))
FIRST_DATA_ROW = 2
# STT is a plain counter in the list, so it should land as a number rather than as text.
NUMERIC_EXPORT_HEADERS = NUMERIC_HEADERS | {"STT"}


def export_cell_value(header, raw):
    raw = raw.strip()
    if not raw:
        return None
    if header not in NUMERIC_EXPORT_HEADERS:
        return raw
    try:
        number = float(raw)
    except ValueError:
        return raw
    return int(number) if number.is_integer() else number


def build_list_workbook(records, template_path):
    """Return the template filled with `records`, plus "Khu phố" and "Diện" columns appended at the end."""
    workbook = load_workbook(template_path)
    sheet = workbook.active
    headers = {}
    for cell in sheet[1]:
        header = normalized_header(cell.value)
        if header and header not in headers:
            headers[header] = cell.column
    if not headers:
        raise WorkbookError("Hàng 1 của tệp mẫu danh sách không có tiêu đề cột nào.")

    for app_header, _ in APP_COLUMNS:
        if app_header in headers:
            continue
        column = max(headers.values()) + 1
        heading = sheet.cell(1, column)
        heading.value = app_header
        neighbour = sheet.cell(1, column - 1)
        if neighbour.has_style:
            heading._style = copy.copy(neighbour._style)
        sheet.column_dimensions[heading.column_letter].width = 26
        headers[app_header] = column

    # The template ships with sample people whose rows may carry one-off emphasis (the first sample
    # was bold). Take the look of the first blank ruled row instead, so every exported row matches.
    last_template_row = sheet.max_row
    style_row = FIRST_DATA_ROW
    for row in range(FIRST_DATA_ROW, last_template_row + 1):
        if all(sheet.cell(row, column).value is None for column in range(1, sheet.max_column + 1)):
            style_row = row
            break

    for row in range(FIRST_DATA_ROW, last_template_row + 1):
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row, column).value = None
        # Sample heights were sized for their long text; let Excel size the real rows.
        sheet.row_dimensions[row].height = None

    app_keys = dict(APP_COLUMNS)
    for offset, record in enumerate(records):
        row = FIRST_DATA_ROW + offset
        if row != style_row:
            copy_row_style(sheet, style_row, row)
        data = record["data"]
        for header, column in headers.items():
            cell = sheet.cell(row, column)
            if header in app_keys:
                cell.value = record.get(app_keys[header]) or None
                neighbour = sheet.cell(row, column - 1)
                if neighbour.has_style:
                    cell._style = copy.copy(neighbour._style)
            else:
                cell.value = export_cell_value(header, data.get(header, ""))

    # Leave no trailing ruled-but-empty rows from the template.
    first_unused = FIRST_DATA_ROW + len(records)
    if first_unused <= last_template_row:
        sheet.delete_rows(first_unused, last_template_row - first_unused + 1)
    return workbook
