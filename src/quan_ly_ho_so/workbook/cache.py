"""Workbook path validation and the SQLite read/search cache."""

import json
import logging
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from quan_ly_ho_so.config import CACHE_SEARCH_SCOPE, DELETED_AT_HEADER, FILTER_DB_COLUMNS, SUPPORTED_SUFFIXES
from quan_ly_ho_so.errors import StaleWorkbookError, WorkbookError
from quan_ly_ho_so.state import current_database_path, state
from quan_ly_ho_so.utils.text import display_value, fold, normalized_header


def file_signature(path):
    stat = path.stat()
    return (stat.st_mtime_ns, stat.st_size)


def signature_token(signature):
    if not signature:
        return ""
    return f"{signature[0]}:{signature[1]}"


def workbook_lock_path(path):
    return path.with_name(f"~${path.name}")


def workbook_appears_open(path):
    lock = workbook_lock_path(path)
    if not lock.is_file():
        return False
    try:
        age = datetime.now().timestamp() - lock.stat().st_mtime
    except OSError:
        return False
    # Excel lock files can be left behind after a crash. Treat only a current lock as active.
    return age < 2 * 60 * 60


def validate_workbook_path(path):
    path = Path(path).expanduser().resolve()
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise WorkbookError("Hãy chọn tệp Excel có đuôi .xlsx hoặc .xlsm.")
    if not path.is_file():
        raise WorkbookError("Không tìm thấy tệp Excel.")
    return path


def database_connection():
    path = current_database_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS cache_metadata (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS records (
            id INTEGER PRIMARY KEY,
            row_number INTEGER NOT NULL,
            stt TEXT NOT NULL,
            name TEXT NOT NULL,
            birth_year TEXT NOT NULL,
            citizen_id TEXT NOT NULL,
            occupation TEXT NOT NULL,
            education TEXT NOT NULL,
            ethnicity TEXT NOT NULL,
            religion TEXT NOT NULL,
            address TEXT NOT NULL,
            search_text TEXT NOT NULL,
            birth_year_key TEXT NOT NULL,
            occupation_key TEXT NOT NULL,
            education_key TEXT NOT NULL,
            ethnicity_key TEXT NOT NULL,
            religion_key TEXT NOT NULL,
            address_key TEXT NOT NULL,
            citizen_id_key TEXT NOT NULL,
            data_json TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS records_stt_idx ON records(stt);
        CREATE INDEX IF NOT EXISTS records_birth_year_idx ON records(birth_year_key);
        CREATE INDEX IF NOT EXISTS records_occupation_idx ON records(occupation_key);
        CREATE INDEX IF NOT EXISTS records_education_idx ON records(education_key);
        CREATE INDEX IF NOT EXISTS records_ethnicity_idx ON records(ethnicity_key);
        CREATE INDEX IF NOT EXISTS records_religion_idx ON records(religion_key);
        CREATE INDEX IF NOT EXISTS records_citizen_id_idx ON records(citizen_id_key);
        """
    )
    connection.commit()
    return connection


def metadata_values(connection):
    return {row["key"]: row["value"] for row in connection.execute("SELECT key, value FROM cache_metadata")}


def cached_workbook_info(path, signature):
    try:
        connection = database_connection()
        try:
            metadata = metadata_values(connection)
        finally:
            connection.close()
        if (
            metadata.get("workbook") != str(path)
            or metadata.get("signature") != signature_token(signature)
            or metadata.get("search_scope") != CACHE_SEARCH_SCOPE
        ):
            return None
        headers = json.loads(metadata["headers"])
        return metadata["sheet"], headers, int(metadata["record_count"]), datetime.fromisoformat(metadata["loaded_at"])
    except (OSError, sqlite3.Error, KeyError, TypeError, ValueError, json.JSONDecodeError) as error:
        logging.warning("Could not reuse SQLite cache: %s", error)
        return None


def sync_workbook_to_database(path):
    """Stream one workbook snapshot into SQLite without retaining all rows in RAM."""
    path = validate_workbook_path(path)
    signature_before = file_signature(path)
    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=keep_vba)
    connection = None
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = {}
        for column, cell in enumerate(first_row, start=1):
            header = normalized_header(cell.value)
            if header and header not in headers:
                headers[header] = column
        if "STT" not in headers or "TÊN THƯỜNG DÙNG" not in headers:
            raise WorkbookError("Hàng 1 của tệp Excel phải có cột STT và TÊN THƯỜNG DÙNG.")

        deleted_at_column = headers.get(normalized_header(DELETED_AT_HEADER))

        connection = database_connection()
        insert_sql = """
            INSERT INTO records (
                row_number, stt, name, birth_year, citizen_id, occupation, education, ethnicity, religion,
                address, search_text, birth_year_key, occupation_key, education_key, ethnicity_key,
                religion_key, address_key, citizen_id_key, data_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        count = 0
        batch = []
        loaded_at = datetime.now()
        with connection:
            connection.execute("DELETE FROM records")
            for row_number, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
                data = {
                    header: display_value(values[column - 1] if column <= len(values) else None)
                    for header, column in headers.items()
                }
                stt = data.get("STT", "")
                name = data.get("TÊN THƯỜNG DÙNG", "")
                if not stt or not name:
                    continue
                if deleted_at_column and deleted_at_column <= len(values):
                    deleted_at = display_value(values[deleted_at_column - 1])
                    if deleted_at:
                        continue  # Skip soft-deleted records
                birth_year = data.get("Năm sinh", "")
                citizen_id = data.get("Căn cước", "")
                occupation = data.get("Nghề nghiệp", "")
                education = data.get("Văn hóa", "")
                ethnicity = data.get("Dân tộc", "")
                religion = data.get("Tôn giáo", "")
                address = f"{data.get('Thường trú', '')} {data.get('Nơi ở hiện nay', '')}".strip()
                # Quick search intentionally matches only the identity fields advertised by the UI.
                # Advanced filters cover occupation, education, address, and the remaining profile data.
                searchable = " ".join([stt, name, citizen_id])
                batch.append((
                    row_number, stt, name, birth_year, citizen_id, occupation, education, ethnicity, religion,
                    address, fold(searchable), fold(birth_year), fold(occupation), fold(education), fold(ethnicity),
                    fold(religion), fold(address), fold(citizen_id),
                    json.dumps(data, ensure_ascii=False, separators=(",", ":")),
                ))
                count += 1
                if len(batch) >= 500:
                    connection.executemany(insert_sql, batch)
                    batch.clear()
            if batch:
                connection.executemany(insert_sql, batch)
            signature_after = file_signature(path)
            if signature_after != signature_before:
                raise StaleWorkbookError("Tệp Excel thay đổi trong lúc đồng bộ. Hãy thử làm mới lại.")
            metadata = {
                "workbook": str(path),
                "signature": signature_token(signature_after),
                "sheet": sheet.title,
                "headers": json.dumps(headers, ensure_ascii=False, separators=(",", ":")),
                "record_count": str(count),
                "loaded_at": loaded_at.isoformat(),
                "search_scope": CACHE_SEARCH_SCOPE,
            }
            connection.executemany(
                "INSERT OR REPLACE INTO cache_metadata(key, value) VALUES (?, ?)", metadata.items()
            )
        return sheet.title, headers, count, signature_after, loaded_at
    finally:
        workbook.close()
        if connection is not None:
            connection.close()


def database_record_count():
    connection = database_connection()
    try:
        return int(connection.execute("SELECT COUNT(*) FROM records").fetchone()[0])
    finally:
        connection.close()


def row_to_record(row):
    return {
        "row": row["row_number"],
        "stt": row["stt"],
        "name": row["name"],
        "birth_year": row["birth_year"],
        "citizen_id": row["citizen_id"],
        "occupation": row["occupation"],
        "data": json.loads(row["data_json"]),
    }


def find_record(stt):
    connection = database_connection()
    try:
        row = connection.execute(
            "SELECT * FROM records WHERE stt = ? ORDER BY row_number LIMIT 1", (display_value(stt),)
        ).fetchone()
        return row_to_record(row) if row is not None else None
    finally:
        connection.close()


def find_record_by_citizen_id(citizen_id):
    """Look up a record by citizen ID (CCCD)."""
    citizen_id = display_value(citizen_id).strip()
    if not citizen_id:
        return None
    connection = database_connection()
    try:
        row = connection.execute(
            "SELECT * FROM records WHERE citizen_id_key = ? ORDER BY row_number LIMIT 1",
            (fold(citizen_id),),
        ).fetchone()
        return row_to_record(row) if row is not None else None
    finally:
        connection.close()


def escaped_like(value):
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def query_records(query, filters, limit, offset):
    conditions = []
    parameters = []
    query_key = fold(query)
    if query_key:
        conditions.append("search_text LIKE ? ESCAPE '\\'")
        parameters.append(f"%{escaped_like(query_key)}%")
    for key, (_, key_column) in FILTER_DB_COLUMNS.items():
        value = fold(filters.get(key, ""))
        if value:
            conditions.append(f"{key_column} = ?")
            parameters.append(value)
    address = fold(filters.get("address", ""))
    if address:
        conditions.append("address_key LIKE ? ESCAPE '\\'")
        parameters.append(f"%{escaped_like(address)}%")
    where = f" WHERE {' AND '.join(conditions)}" if conditions else ""

    connection = database_connection()
    try:
        count = int(connection.execute(f"SELECT COUNT(*) FROM records{where}", parameters).fetchone()[0])
        rows = connection.execute(
            f"SELECT * FROM records{where} ORDER BY row_number LIMIT ? OFFSET ?",
            [*parameters, limit, offset],
        ).fetchall()
        return [row_to_record(row) for row in rows], count
    finally:
        connection.close()


def query_filter_options(key):
    if key not in FILTER_DB_COLUMNS:
        raise ValueError(f"Unknown filter: {key}")
    value_column, key_column = FILTER_DB_COLUMNS[key]
    connection = database_connection()
    try:
        rows = connection.execute(
            f"""
            SELECT {value_column} AS value, {key_column} AS normalized
            FROM records
            WHERE {key_column} <> ''
            ORDER BY row_number
            """
        )
        values = {}
        for row in rows:
            values.setdefault(row["normalized"], row["value"])
        return sorted(values.values(), key=fold)
    finally:
        connection.close()
