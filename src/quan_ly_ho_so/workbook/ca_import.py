"""Read the Ministry of Public Security's per-khu-phố lists and turn their new people into workbook rows."""

import io
import re

from openpyxl import load_workbook

from quan_ly_ho_so.config import QUARTER_CHOICES
from quan_ly_ho_so.forms.fields import coerce_cell_value
from quan_ly_ho_so.utils.text import display_value, fold, normalized_header
from quan_ly_ho_so.workbook.dien_export import (
    ADDRESS, BACKGROUND, EDUCATION, FAMILY, JOB, NAME_AND_BIRTH, NOTE, classify_column,
)

CITIZEN_ID = "citizen_id"
QUARTER_NEW = "quarter_new"
# The workbook's catch-all column, where anything that cannot be placed safely is kept as written.
OTHER_INFO = "TẤT CẢ THÔNG TIN KHÁC\nLIÊN QUAN ĐẾN THANH NIÊN"
NOTE_HEADER = "GHI CHÚ\n(Về Sức khỏe; gia cảnh)"
# Father, mother, spouse: the order the lists print them in.
FAMILY_HEADERS = (
    ("Tên cha", "năm sinh cha", "Nghề nghiệp cha"),
    ("Tên mẹ", "Năm sinh mẹ", "Nghề nghiệp mẹ"),
    ("Tên vợ", "năm sinh vợ", "nghề nghiệp vợ"),
)


def column_kind(header):
    key = " ".join(fold(header).split())
    if key == "cccd" or "can cuoc" in key:
        return CITIZEN_ID
    if "khu pho moi" in key:
        return QUARTER_NEW
    if "khu pho cu" in key:
        return None
    return classify_column(header)


def lines_of(value):
    return [line.strip() for line in display_value(value).replace("\r", "").split("\n") if line.strip()]


def normalized_citizen_id(value):
    digits = re.sub(r"\s+", "", display_value(value))
    # A leading zero is lost when Excel keeps the number as a number.
    return digits.zfill(12) if digits.isdigit() and len(digits) == 11 else digits


def _birth(lines, notes):
    """Split "Name\\ndd/mm/yyyy" into the name and the date's three numbers."""
    name = lines[0] if lines else ""
    values = {}
    rest = " ".join(lines[1:])
    full = re.search(r"(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{4})", rest)
    if full:
        values["Ngày sinh"], values["Tháng sinh"], values["Năm sinh"] = full.groups()
    elif re.fullmatch(r"\d{4}", rest.strip()):
        values["Năm sinh"] = rest.strip()
    elif rest.strip():
        notes.append(f"không đọc được ngày sinh \"{rest.strip()}\"")
    return name, values


def _name_and_year(line):
    match = re.match(r"^(.*?)[\s,;]+(\d{4})[\s,;]*$", line)
    return (match.group(1).strip(), match.group(2)) if match else (line.strip(" ,;"), "")


def _family(lines, values, notes):
    # Pairs of "Name year," and a job line. A list of any other length cannot be told apart
    # (which parent is it?), so it is kept as written rather than guessed at.
    if len(lines) not in (4, 6):
        if lines:
            values[OTHER_INFO] = "Gia đình (theo file CA):\n" + "\n".join(lines)
            notes.append("thông tin gia đình để ở cột \"Tất cả thông tin khác\"")
        return
    for (name_header, year_header, job_header), start in zip(FAMILY_HEADERS, range(0, len(lines), 2)):
        name, year = _name_and_year(lines[start])
        values[name_header], values[year_header], values[job_header] = name, year, lines[start + 1]


def parse_row(row, columns):
    """One list row -> {"name", "citizen_id", "quarter", "values", "notes"}, or None for a banner/blank row."""
    def cell(kind):
        index = columns.get(kind)
        return row[index] if index is not None and index < len(row) else None

    notes = []
    name, values = _birth(lines_of(cell(NAME_AND_BIRTH)), notes)
    if not name:
        return None
    values["Tên khai sinh"] = name
    values["TÊN THƯỜNG DÙNG"] = name.upper()
    citizen_id = normalized_citizen_id(cell(CITIZEN_ID))
    values["Căn cước"] = citizen_id

    job = lines_of(cell(JOB))
    values["Nghề nghiệp"] = job[0] if job else ""
    values["Nơi làm việc, học tập"] = "; ".join(job[1:])
    address = lines_of(cell(ADDRESS))
    values["Thường trú"] = address[0] if address else ""
    values["Nơi ở hiện nay"] = "; ".join(address[1:])
    background = lines_of(cell(BACKGROUND))
    values["Dân tộc"] = background[0] if background else ""
    values["Tôn giáo"] = background[1] if len(background) > 1 else ""
    education = lines_of(cell(EDUCATION))
    for header, text in zip(("Văn hóa", "Chuyên môn", "Ngành đào tạo"), education):
        values[header] = text
    _family(lines_of(cell(FAMILY)), values, notes)
    values[NOTE_HEADER] = "\n".join(lines_of(cell(NOTE)))

    quarter = ""
    raw_quarter = display_value(cell(QUARTER_NEW))
    if raw_quarter:
        quarter = f"Khu phố {raw_quarter}"
        if quarter not in QUARTER_CHOICES:
            notes.append(f"khu phố \"{raw_quarter}\" không có trong danh sách, chưa gán khu phố")
            quarter = ""
    return {
        "name": name,
        "citizen_id": citizen_id,
        "quarter": quarter,
        "values": {header: coerce_cell_value(header, str(value)) for header, value in values.items()},
        "notes": notes,
    }


def read_ca_file(data):
    """All people in one CA workbook, from every sheet laid out like the lists."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    people = []
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            for header_index, header_row in enumerate(rows[:5]):
                columns = {}
                for index, header in enumerate(header_row):
                    kind = column_kind(normalized_header(header)) if header is not None else None
                    if kind is not None and kind not in columns:
                        columns[kind] = index
                if CITIZEN_ID in columns and NAME_AND_BIRTH in columns:
                    people.extend(person for person in (parse_row(row, columns) for row in rows[header_index + 1:]) if person)
                    break
    finally:
        workbook.close()
    return people


def read_ca_files(files):
    """files: [(name, bytes)] -> (people, errors). Each person also records the file it came from."""
    people, errors = [], []
    for name, data in files:
        try:
            found = read_ca_file(data)
        except Exception as error:
            errors.append(f"{name}: không đọc được tệp ({error})")
            continue
        if not found:
            errors.append(f"{name}: không thấy bảng danh sách có cột CCCD và họ tên khai sinh")
        for person in found:
            person["file"] = name
        people.extend(found)
    return people, errors


def plan_import(people, known_citizen_ids):
    """Sort people into those to add and those to leave alone, in the order they were read."""
    plan = {"add": [], "existing": [], "no_citizen_id": [], "repeated": []}
    seen = set()
    for person in people:
        citizen_id = person["citizen_id"]
        if not citizen_id:
            plan["no_citizen_id"].append(person)
        elif citizen_id in known_citizen_ids:
            plan["existing"].append(person)
        elif citizen_id in seen:
            plan["repeated"].append(person)
        else:
            seen.add(citizen_id)
            plan["add"].append(person)
    return plan
