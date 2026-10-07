"""Fill the per-diện Excel templates, one file for each diện the manager picked."""

import copy
import re

from openpyxl import load_workbook

from quan_ly_ho_so.errors import WorkbookError
from quan_ly_ho_so.utils.text import fold, normalized_header
from quan_ly_ho_so.workbook.writer import copy_row_style

# Each diện is reported on its own official form.
DIEN_TEMPLATES = {
    "MIEN_DANG_KY": "Mau_Mien_Dang_Ky.xlsx",
    "MIEN_GOI_NHAP_NGU": "Mau_Mien_Goi_Nhap_Ngu.xlsx",
    "TAM_HOAN_SUC_KHOE": "Mau_Suc_Khoe.xlsx",
    "TAM_HOAN_HSSV": "Mau_HSSV.xlsx",
    "TAM_HOAN_CHINH_SACH": "Mau_Tam_Hoan_Chinh_Sach.xlsx",
    "TAM_HOAN_DAN_QUAN_THUONG_TRUC": "Mau_Dan_Quan_Thuong_Truc.xlsx",
    "TAM_HOAN_HOC_VAN_THAP": "Mau_Hoc_Van_Thap.xlsx",
    "KHONG_TUYEN_CHON": "Mau_Khong_Tuyen_Chon.xlsx",
    "CHUA_XET_TUYEN": "Mau_Chua_Xet_Tuyen.xlsx",
    "DU_DIEU_KIEN": "Mau_Du_Dieu_Kien.xlsx",
    "DAN_QUAN": "Mau_Dan_Quan.xlsx",
    "DU_HOC": "Mau_Du_Hoc.xlsx",
    "GIA_CANH": "Mau_Gia_Canh.xlsx",
    "LAO_DONG_NUOC_NGOAI": "Mau_LDNN.xlsx",
    "TON_GIAO": "Mau_Ton_Giao.xlsx",
    "VANG_MAT_DIA_PHUONG": "Mau_Vang_Mat_Dia_Phuong.xlsx",
}

MASTER_STT = "master_stt"
DIEN_INDEX = "dien_index"
NAME_AND_BIRTH = "name_and_birth"
JOB = "job"
ADDRESS = "address"
BACKGROUND = "background"
EDUCATION = "education"
FAMILY = "family"
NOTE = "note"
QUARTER = "quarter"


def classify_column(header):
    """Name the kind of data a template column wants, from its (multi-line) heading."""
    key = " ".join(fold(header).split())
    if not key:
        return None
    # "Họ tên cha; năm sinh, nghề nghiệp" also mentions a job, so claim it before the job column.
    if "ho ten cha" in key:
        return FAMILY
    if key.startswith("stt theo dien"):
        return DIEN_INDEX
    if key in ("so tt", "stt") or key.startswith("so tt ") or key.startswith("stt "):
        return MASTER_STT
    if "ten khai sinh" in key:
        return NAME_AND_BIRTH
    if "thuong tru" in key:
        return ADDRESS
    if "hoc van" in key:
        return EDUCATION
    if "dan toc" in key:
        return BACKGROUND
    if "nghe nghiep" in key:
        return JOB
    if "khu" in key and "pho" in key:
        return QUARTER
    if "ghi chu" in key:
        return NOTE
    # "Lý do cụ thể" and similar have no counterpart in the workbook, so they stay blank.
    return None


def _stack(*parts):
    """Join the lines of one cell, dropping the pieces this record has no data for."""
    return "\n".join(part for part in parts if part) or None


def _birth_date(data):
    day, month, year = data.get("Ngày sinh", ""), data.get("Tháng sinh", ""), data.get("Năm sinh", "")
    if not year:
        return ""
    if day and month:
        return f"{int(float(day)):02d}/{int(float(month)):02d}/{year}"
    return year


# The workbook spells "no spouse yet" out in the name cell; that is an absence, not a person.
ABSENT_PERSON = {"chua co", "khong co", "khong", "chua", "k co", "khong ro"}


def _person(data, name_header, year_header, job_header):
    name = data.get(name_header, "").strip()
    if not name or fold(name) in ABSENT_PERSON:
        return ""
    year = data.get(year_header, "").strip()
    heading = "; ".join(piece for piece in (name, year) if piece) + ";"
    return _stack(heading, data.get(job_header, "").strip())


def _quarter_number(quarter):
    """Templates list the neighbourhood as a bare number, the way their own samples do."""
    match = re.search(r"\d+", quarter or "")
    return int(match.group(0)) if match else (quarter or None)


def cell_content(kind, record, header, position):
    data = record["data"]
    if kind == MASTER_STT:
        stt = record["stt"]
        return int(float(stt)) if stt.replace(".", "").isdigit() else stt
    if kind == DIEN_INDEX:
        return position
    if kind == NAME_AND_BIRTH:
        return _stack(data.get("Tên khai sinh") or data.get("TÊN THƯỜNG DÙNG", ""), _birth_date(data))
    if kind == JOB:
        return _stack(data.get("Nghề nghiệp", ""), data.get("Nơi làm việc, học tập", ""))
    if kind == ADDRESS:
        return _stack(data.get("Thường trú", ""), data.get("Nơi ở hiện nay", ""))
    if kind == BACKGROUND:
        pieces = []
        if "thanh phan" in fold(header):
            pieces = [data.get("Thành phần", ""), data.get("Bản thân", "")]
        return _stack(*pieces, data.get("Dân tộc", ""), data.get("Tôn giáo", ""))
    if kind == EDUCATION:
        return _stack(data.get("Văn hóa", ""), data.get("Chuyên môn", ""), data.get("Ngành đào tạo", ""))
    if kind == FAMILY:
        return _stack(
            _person(data, "Tên cha", "năm sinh cha", "Nghề nghiệp cha"),
            _person(data, "Tên mẹ", "Năm sinh mẹ", "Nghề nghiệp mẹ"),
            _person(data, "Tên vợ", "năm sinh vợ", "nghề nghiệp vợ"),
        )
    if kind == NOTE:
        return _stack(data.get("GHI CHÚ\n(Về Sức khỏe; gia cảnh)", ""))
    if kind == QUARTER:
        return _quarter_number(record.get("quarter"))
    return None


def _layout(sheet):
    """Locate the heading row and the first row of sample people below it."""
    header_row = None
    for row in range(1, min(sheet.max_row, 10) + 1):
        if classify_column(normalized_header(sheet.cell(row, 1).value)) in (MASTER_STT, DIEN_INDEX):
            header_row = row
            break
    if header_row is None:
        raise WorkbookError("Không tìm thấy hàng tiêu đề (cột Số TT) trong tệp mẫu.")

    for row in range(header_row + 1, sheet.max_row + 1):
        filled = sum(1 for column in range(1, sheet.max_column + 1) if sheet.cell(row, column).value is not None)
        # Section banners such as "DU HỌC" span the sheet but fill a single cell; people fill many.
        if filled > 1:
            return header_row, row
    return header_row, header_row + 1


def build_dien_workbook(records, template_path):
    """Return `template_path` filled with `records`, replacing the sample people it ships with."""
    workbook = load_workbook(template_path)
    sheet = workbook.active
    header_row, first_data_row = _layout(sheet)

    columns = {}
    for column in range(1, sheet.max_column + 1):
        header = normalized_header(sheet.cell(header_row, column).value)
        kind = classify_column(header)
        if kind is not None:
            columns[column] = (kind, header)
    # With only one numbering column the form just counts its own rows.
    if not any(kind == DIEN_INDEX for kind, _ in columns.values()):
        columns = {
            column: (DIEN_INDEX if kind == MASTER_STT else kind, header)
            for column, (kind, header) in columns.items()
        }

    last_template_row = sheet.max_row
    for row in range(first_data_row, last_template_row + 1):
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row, column).value = None
        sheet.row_dimensions[row].height = None

    for offset, record in enumerate(records):
        row = first_data_row + offset
        if row != first_data_row:
            copy_row_style(sheet, first_data_row, row)
        for column, (kind, header) in columns.items():
            sheet.cell(row, column).value = cell_content(kind, record, header, offset + 1)

    first_unused = first_data_row + len(records)
    if first_unused <= last_template_row:
        sheet.delete_rows(first_unused, last_template_row - first_unused + 1)
    return workbook
