#!/usr/bin/env python3
"""Generate one youth profile Word document from a selected Excel row."""

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_filename(value):
    name = re.sub(r'[\\/:*?"<>|]+', "-", value).strip(" .")
    return name or "unnamed-record"


def replace_tokens(document, values):
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            # Reassigning run.text rebuilds the run from plain text, which would
            # silently delete non-text content (like the photo-box drawing) from
            # runs that don't hold a token. Only touch runs that need replacing.
            if "{{" not in run.text:
                continue
            new_text = run.text
            for key, value in values.items():
                new_text = new_text.replace("{{" + key + "}}", value)
            if new_text != run.text:
                run.text = new_text


def unique_output_path(output_dir, name):
    output = output_dir / (safe_filename(name) + ".docx")
    if not output.exists():
        return output
    number = 2
    while True:
        candidate = output_dir / f"{safe_filename(name)} ({number}).docx"
        if not candidate.exists():
            return candidate
        number += 1


def build_document(workbook_path, sheet_name, row_number, template_path):
    if row_number < 3:
        raise ValueError("Select a person row (row 3 or later), not a header row.")
    if not workbook_path.is_file() or not template_path.is_file():
        raise FileNotFoundError("Workbook or Word template was not found.")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_vba=workbook_path.suffix.lower() == ".xlsm")
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet was not found: {sheet_name}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        headers = [text(cell.value) for cell in sheet[1]]
        row = [text(cell.value) for cell in sheet[row_number]]
        data = dict(zip(headers, row))

        def field(header):
            return data.get(header, "")

        birth_name = field("Tên khai sinh") or field("TÊN THƯỜNG DÙNG")
        siblings = [field(f"Anh chị em {number}") for number in range(1, 8)]
        values = {
            "record_stt": field("STT"),
            "birth_name_upper": birth_name.upper(),
            "common_name": field("TÊN THƯỜNG DÙNG") or birth_name,
            "birth_day": field("Ngày sinh"),
            "birth_month": field("Tháng sinh"),
            "birth_year": field("Năm sinh"),
            "gender": "", "citizen_id": field("Căn cước"),
            "birthplace": field("Nơi đăng ký khai sinh (2 cấp)"),
            "hometown": field("Quê quán (2 cấp)"),
            "ethnicity": field("Dân tộc"), "religion": field("Tôn giáo"),
            "nationality": "Việt Nam", "permanent_address": field("Thường trú"),
            "current_address": field("Nơi ở hiện nay"),
            "family_background": field("Thành phần"), "personal_status": field("Bản thân"),
            "education": field("Văn hóa"), "specialization": field("Chuyên môn"),
            "foreign_language": "", "training_major": field("Ngành đào tạo"),
            "occupation": field("Nghề nghiệp"), "salary": "", "grade": "", "rank": "",
            "workplace": field("Nơi làm việc, học tập"),
            "father_name": field("Tên cha"), "father_life_status": field("Sống / Chết"),
            "father_birth_day": field("Ngày sinh cha"), "father_birth_month": field("tháng sinh cha"),
            "father_birth_year": field("năm sinh cha"), "father_occupation": field("Nghề nghiệp cha"),
            "mother_name": field("Tên mẹ"), "mother_life_status": field("Sống / chết (mẹ)"),
            "mother_birth_day": field("Ngày sinh mẹ"), "mother_birth_month": field("Tháng sinh mẹ"),
            "mother_birth_year": field("Năm sinh mẹ"), "mother_occupation": field("Nghề nghiệp mẹ"),
            "spouse_name": field("Tên vợ") or "Chưa có", "spouse_day": "", "spouse_month": "",
            "spouse_year": field("năm sinh vợ"), "sibling_count": field("cha mẹ có bao nhiêu con"),
            "brother_count": field("mấy trai"), "sister_count": field("mấy gái"), "birth_order": field("là Con thứ"),
            "spouse_occupation": field("nghề nghiệp vợ"), "spouse_children": field("Số người con"),
            "father_before_1975": field("Trước 30-4 (cha)\nghi 3 cấp và 2 cấp"),
            "father_after_1975": field("sau 30-4 (cha)\nghi 2 cấp"),
            "father_current": field("hiện nay (cha)\nghi 2 cấp"),
            "mother_before_1975": field("Trước 30-4 (mẹ)\nghi 3 cấp và 2 cấp"),
            "mother_after_1975": field("sau 30-4 (mẹ)\nghi 2 cấp"),
            "mother_current": field("hiện nay (mẹ)\nghi 2 cấp"),
            "siblings": "\n".join(sibling for sibling in siblings if sibling),
            "education_level_1": field("Cấp 1"), "education_level_2": field("Cấp 2"),
            "education_level_3": field("Cấp 3"), "current_history": field("Hiện nay"),
        }
    finally:
        workbook.close()

    document = Document(template_path)
    replace_tokens(document, values)
    return document, values


def generate_document(workbook_path, sheet_name, row_number, template_path, output_dir):
    document, values = build_document(workbook_path, sheet_name, row_number, template_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_name = f"{values['record_stt']} - {values['common_name']}" if values["record_stt"] else values["common_name"]
    output = unique_output_path(output_dir, output_name)
    document.save(output)
    return output


def _boundary_paragraph(document):
    """Return the paragraph element that ends Part A (I-IV) and carries the
    section break into Part B (V-VI). The template always has exactly one
    such mid-body section break."""
    for paragraph in document.paragraphs:
        if paragraph._p.find(f".//{WORD_NS}sectPr") is not None:
            return paragraph._p
    raise ValueError("Template is missing the section break between Part A (I-IV) and Part B (V-VI).")


def split_parts(document):
    """Split a generated document into Part A (header..IV, ending in the
    boundary paragraph that holds Section A's page/column setup) and Part B
    (V-VI, everything after the boundary paragraph up to the document's own
    trailing section properties)."""
    body = document.element.body
    boundary = _boundary_paragraph(document)
    children = list(body)
    boundary_index = children.index(boundary)
    part_a = children[: boundary_index + 1]
    part_b = children[boundary_index + 1 : -1]
    section_properties = boundary.find(f".//{WORD_NS}sectPr")
    return part_a, part_b, section_properties


def combine_booklet(documents, template_path):
    """Combine several generated documents into one printable booklet.

    Printing these forms back-to-back on A3 and folding them into a booklet
    leaves the second column of every person's Part A (I-IV) mostly blank.
    To avoid wasting that space, each sheet instead carries the *previous*
    person's Part B (V-VI) next to the *current* person's Part A, with the
    very last person's Part B wrapped around to the front of the file.
    """
    if not documents:
        raise ValueError("No documents to combine.")

    parts = [split_parts(document) for document in documents]
    count = len(documents)

    combined = Document(template_path)
    body = combined.element.body
    for child in list(body):
        body.remove(child)

    for index in range(count):
        previous_part_b = parts[index - 1][1]
        current_part_a, _, current_section_properties = parts[index]
        for element in previous_part_b:
            body.append(copy.deepcopy(element))
        if index < count - 1:
            for element in current_part_a:
                body.append(copy.deepcopy(element))
        else:
            for element in current_part_a[:-1]:
                body.append(copy.deepcopy(element))
            body.append(copy.deepcopy(current_section_properties))

    return combined


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--sheet", help="Worksheet containing the source data")
    parser.add_argument("--row", required=True, type=int)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    print(generate_document(args.workbook, args.sheet, args.row, args.template, args.output_dir))


if __name__ == "__main__":
    main()
