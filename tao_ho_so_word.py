#!/usr/bin/env python3
"""Generate one youth profile Word document from a selected Excel row."""

import argparse
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_filename(value):
    name = re.sub(r'[\\/:*?"<>|]+', "-", value).strip(" .")
    return name or "unnamed-record"


def replace_tokens(document, values):
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            for key, value in values.items():
                run.text = run.text.replace("{{" + key + "}}", value)


def unique_output_path(output_dir, name):
    output = output_dir / (safe_filename(name) + ".docx")
    if not output.exists():
        return output
    number = 2
    while True:
        candidate = output_dir / f"{safe_filename(name)} ({number}).docx"
        if not candidate.exists():
            return candidate
        number += 1


def generate_document(workbook_path, sheet_name, row_number, template_path, output_dir):
    if row_number < 3:
        raise ValueError("Select a person row (row 3 or later), not a header row.")
    if not workbook_path.is_file() or not template_path.is_file():
        raise FileNotFoundError("Workbook or Word template was not found.")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_vba=workbook_path.suffix.lower() == ".xlsm")
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet was not found: {sheet_name}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        headers = [text(cell.value) for cell in sheet[1]]
        row = [text(cell.value) for cell in sheet[row_number]]
        data = dict(zip(headers, row))

        def field(header):
            return data.get(header, "")

        birth_name = field("Tên khai sinh") or field("TÊN THƯỜNG DÙNG")
        siblings = [field(f"Anh chị em {number}") for number in range(1, 8)]
        values = {
            "record_stt": field("STT"),
            "birth_name_upper": birth_name.upper(),
            "common_name": field("TÊN THƯỜNG DÙNG") or birth_name,
            "birth_day": field("Ngày sinh"),
            "birth_month": field("Tháng sinh"),
            "birth_year": field("Năm sinh"),
            "gender": "", "citizen_id": field("Căn cước"),
            "birthplace": field("Nơi đăng ký khai sinh (2 cấp)"),
            "hometown": field("Quê quán (2 cấp)"),
            "ethnicity": field("Dân tộc"), "religion": field("Tôn giáo"),
            "nationality": "Việt Nam", "permanent_address": field("Thường trú"),
            "current_address": field("Nơi ở hiện nay"),
            "family_background": field("Thành phần"), "personal_status": field("Bản thân"),
            "education": field("Văn hóa"), "specialization": field("Chuyên môn"),
            "foreign_language": "", "training_major": field("Ngành đào tạo"),
            "occupation": field("Nghề nghiệp"), "salary": "", "grade": "", "rank": "",
            "workplace": field("Nơi làm việc, học tập"),
            "father_name": field("Tên cha"), "father_life_status": field("Sống / Chết"),
            "father_birth_day": field("Ngày sinh cha"), "father_birth_month": field("tháng sinh cha"),
            "father_birth_year": field("năm sinh cha"), "father_occupation": field("Nghề nghiệp cha"),
            "mother_name": field("Tên mẹ"), "mother_life_status": field("Sống / chết (mẹ)"),
            "mother_birth_day": field("Ngày sinh mẹ"), "mother_birth_month": field("Tháng sinh mẹ"),
            "mother_birth_year": field("Năm sinh mẹ"), "mother_occupation": field("Nghề nghiệp mẹ"),
            "spouse_name": field("Tên vợ") or "Chưa có", "spouse_day": "", "spouse_month": "",
            "spouse_year": field("năm sinh vợ"), "sibling_count": field("cha mẹ có bao nhiêu con"),
            "brother_count": field("mấy trai"), "sister_count": field("mấy gái"), "birth_order": field("là Con thứ"),
            "spouse_occupation": field("nghề nghiệp vợ"), "spouse_children": field("Số người con"),
            "father_before_1975": field("Trước 30-4 (cha)\nghi 3 cấp và 2 cấp"),
            "father_after_1975": field("sau 30-4 (cha)\nghi 2 cấp"),
            "father_current": field("hiện nay (cha)\nghi 2 cấp"),
            "mother_before_1975": field("Trước 30-4 (mẹ)\nghi 3 cấp và 2 cấp"),
            "mother_after_1975": field("sau 30-4 (mẹ)\nghi 2 cấp"),
            "mother_current": field("hiện nay (mẹ)\nghi 2 cấp"),
            "siblings": "\n".join(sibling for sibling in siblings if sibling),
            "education_level_1": field("Cấp 1"), "education_level_2": field("Cấp 2"),
            "education_level_3": field("Cấp 3"), "current_history": field("Hiện nay"),
        }
    finally:
        workbook.close()

    document = Document(template_path)
    replace_tokens(document, values)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_name = f"{values['record_stt']} - {values['common_name']}" if values["record_stt"] else values["common_name"]
    output = unique_output_path(output_dir, output_name)
    document.save(output)
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--sheet", help="Worksheet containing the source data")
    parser.add_argument("--row", required=True, type=int)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    print(generate_document(args.workbook, args.sheet, args.row, args.template, args.output_dir))


if __name__ == "__main__":
    main()
