import re
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook

import quan_ly_ho_so.app as app
from quan_ly_ho_so.auth import create_user, set_auth_database_path
from quan_ly_ho_so.audit import get_audit_logs
from quan_ly_ho_so.workbook.cache import find_record, find_record_by_citizen_id


def create_sample_workbook(path):
    headers = ["STT"]
    for _, section_headers in app.FORM_GROUPS:
        for header in section_headers:
            if header not in headers:
                headers.append(header)
    records = [
        {
            "STT": 1,
            "TÊN THƯỜNG DÙNG": "NGUYỄN VĂN AN",
            "Năm sinh": 2005,
            "Căn cước": "001205000001",
            "Nghề nghiệp": "Sinh viên",
        },
        {
            "STT": 2,
            "TÊN THƯỜNG DÙNG": "TRẦN THỊ BÌNH",
            "Năm sinh": 2004,
            "Căn cước": "001204000002",
            "Nghề nghiệp": "Nông dân",
        },
    ]
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Hồ sơ"
    sheet.append(headers)
    sheet.append([])
    for record in records:
        sheet.append([record.get(header) for header in headers])
    wb.save(path)
    wb.close()


class NewFeaturesTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "records.xlsx"
        self.previous_database = app.state["database"]
        app.state["database"] = Path(self.temp_dir.name) / "cache.sqlite3"
        set_auth_database_path(Path(self.temp_dir.name) / "users.db")
        create_sample_workbook(self.workbook_path)
        app.set_workbook(self.workbook_path, persist=False)

    def tearDown(self):
        set_auth_database_path(None)
        app.state.update(
            workbook=None,
            sheet=None,
            headers={},
            database=self.previous_database,
            record_count=0,
            signature=None,
            loaded_at=None,
            error=None,
            last_output_dir=None,
        )
        self.temp_dir.cleanup()

    def _login(self, client, username="admin", password="admin123"):
        login_page = client.get("/login")
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', login_page.get_data(as_text=True)).group(1)
        client.post("/login", data={"username": username, "password": password, "csrf_token": csrf})

    def test_soft_delete_preserves_row_in_excel_and_skips_in_cache(self):
        client = app.app.test_client()
        self._login(client)

        home = client.get("/")
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', home.get_data(as_text=True)).group(1)
        sig = re.search(r'name="signature" value="([^"]*)"', home.get_data(as_text=True)).group(1)

        # Delete record 1
        resp = client.post("/person/1/delete", data={"csrf_token": csrf, "signature": sig}, follow_redirects=True)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Đã xóa hồ sơ STT 1 (xóa mềm)", resp.get_data(as_text=True))

        # Record 1 should not appear in active records
        self.assertEqual(app.state["record_count"], 1)
        self.assertIsNone(find_record("1"))
        self.assertIsNotNone(find_record("2"))

        # Verify Excel file physically still has the row with deletion timestamp
        wb = load_workbook(self.workbook_path, read_only=True)
        sheet = wb.active
        # Header should include 'Đã xóa lúc'
        header_row = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        self.assertIn("Đã xóa lúc", header_row)
        del_col_idx = header_row.index("Đã xóa lúc")

        # Row 3 (STT 1) still exists in Excel and has deleted_at value
        row3 = [cell.value for cell in list(sheet.iter_rows(min_row=3, max_row=3))[0]]
        self.assertIsNotNone(row3[del_col_idx])
        self.assertNotEqual(str(row3[del_col_idx]).strip(), "")
        wb.close()

    def test_crud_audit_logging(self):
        client = app.app.test_client()
        self._login(client)

        # 1. Update an existing record
        edit_page = client.get("/person/2/edit")
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', edit_page.get_data(as_text=True)).group(1)
        sig = re.search(r'name="signature" value="([^"]*)"', edit_page.get_data(as_text=True)).group(1)

        name_col = app.state["headers"]["TÊN THƯỜNG DÙNG"]
        cccd_col = app.state["headers"]["Căn cước"]
        client.post(
            "/person/save",
            data={
                "csrf_token": csrf,
                "original_stt": "2",
                "signature": sig,
                f"column_{name_col}": "TRẦN THỊ BÌNH (ĐÃ SỬA)",
                f"column_{cccd_col}": "001204000002",
            },
            follow_redirects=True,
        )

        logs, count = get_audit_logs()
        self.assertGreaterEqual(count, 1)
        update_log = next((l for l in logs if l["action"] == "update"), None)
        self.assertIsNotNone(update_log)
        self.assertEqual(update_log["username"], "admin")
        self.assertEqual(update_log["stt"], "2")
        self.assertIn("changes", update_log["details"])
        change_fields = [c["field"] for c in update_log["details"]["changes"]]
        self.assertIn("TÊN THƯỜNG DÙNG", change_fields)

        # 2. Check admin audit log web view
        audit_resp = client.get("/admin/audit-log")
        self.assertEqual(audit_resp.status_code, 200)
        self.assertIn("Nhật ký hoạt động", audit_resp.get_data(as_text=True))
        self.assertIn("Cập nhật", audit_resp.get_data(as_text=True))
        self.assertIn("TRẦN THỊ BÌNH (ĐÃ SỬA)", audit_resp.get_data(as_text=True))

        # 3. Check home page action dropdown button and row-level audit log link
        home_resp = client.get("/")
        home_html = home_resp.get_data(as_text=True)
        self.assertIn('class="button secondary dropdown-toggle"', home_html)
        self.assertIn('/admin/audit-log?stt=2', home_html)

        # 4. Check filtered audit log by STT
        stt_resp = client.get("/admin/audit-log?stt=2")
        self.assertEqual(stt_resp.status_code, 200)
        stt_html = stt_resp.get_data(as_text=True)
        self.assertIn("Lịch sử thao tác cho riêng hồ sơ STT 2", stt_html)
        self.assertIn("TRẦN THỊ BÌNH (ĐÃ SỬA)", stt_html)

        # Filtering for nonexistent STT returns 0 records
        empty_stt_resp = client.get("/admin/audit-log?stt=999999")
        self.assertIn("Chưa có lượt thao tác nào cho hồ sơ STT 999999", empty_stt_resp.get_data(as_text=True))

    def test_guest_lookup_page_unauthenticated(self):
        client = app.app.test_client()

        # 1. Access guest lookup without authentication (GET /tra-cuu)
        resp = client.get("/tra-cuu")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertIn("Tra cứu hồ sơ thanh niên", html)
        self.assertIn("Số căn cước công dân (CCCD)", html)

        csrf = re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)

        # 2. Lookup existing CCCD -> Readonly mode with hover tooltip
        resp_existing = client.post("/tra-cuu", data={"cccd": "001205000001", "csrf_token": csrf})
        self.assertEqual(resp_existing.status_code, 200)
        existing_html = resp_existing.get_data(as_text=True)
        self.assertIn("Hồ sơ đã tồn tại", existing_html)
        self.assertIn("Kiểm tra thông tin và liên hệ admin nếu cần sửa đổi", existing_html)
        self.assertIn("NGUYỄN VĂN AN", existing_html)
        self.assertIn("readonly", existing_html)

        # 3. Lookup non-existing CCCD -> Create mode with form
        resp_new = client.post("/tra-cuu", data={"cccd": "001209999999", "csrf_token": csrf})
        self.assertEqual(resp_new.status_code, 200)
        new_html = resp_new.get_data(as_text=True)
        self.assertIn("Không tìm thấy hồ sơ với CCCD", new_html)
        self.assertIn("001209999999", new_html)
        self.assertIn("Gửi hồ sơ", new_html)

        # 4. Guest submit new record
        sig = re.search(r'name="signature" value="([^"]*)"', new_html).group(1)
        name_col = app.state["headers"]["TÊN THƯỜNG DÙNG"]
        cccd_col = app.state["headers"]["Căn cước"]

        save_resp = client.post(
            "/tra-cuu/save",
            data={
                "csrf_token": csrf,
                "cccd": "001209999999",
                "signature": sig,
                f"column_{name_col}": "HOÀNG GIA BẢO",
                f"column_{cccd_col}": "001209999999",
            },
            follow_redirects=True,
        )
        self.assertEqual(save_resp.status_code, 200)
        save_html = save_resp.get_data(as_text=True)
        self.assertIn("Đã tạo hồ sơ thành công", save_html)
        self.assertIn("HOÀNG GIA BẢO", save_html)
        self.assertIn("Kiểm tra thông tin và liên hệ admin nếu cần sửa đổi", save_html)

        # Verify record exists in cache
        record = find_record_by_citizen_id("001209999999")
        self.assertIsNotNone(record)
        self.assertEqual(record["name"], "HOÀNG GIA BẢO")


if __name__ == "__main__":
    unittest.main()
