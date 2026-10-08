import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import quan_ly_ho_so.app as app
from quan_ly_ho_so.word.export import generate_document


TEMPLATE = Path(__file__).resolve().parent.parent / "Mau_Ho_So_Thanh_Nien.docx"


def create_test_workbook(path):
    headers = ["STT"]
    for _, section_headers in app.FORM_GROUPS:
        for header in section_headers:
            if header not in headers:
                headers.append(header)
    records = [
        {
            "STT": 1,
            "TÊN THƯỜNG DÙNG": "NGUYỄN VĂN AN",
            "Tên khai sinh": "Nguyễn Văn An",
            "Năm sinh": 2005,
            "Căn cước": "001205000001",
            "Nghề nghiệp": "Sinh viên",
            "Văn hóa": "12/12",
            "Dân tộc": "Kinh",
            "Tôn giáo": "Không",
            "Thường trú": "Hà Nội",
        },
        {
            "STT": 2,
            "TÊN THƯỜNG DÙNG": "TRẦN THỊ BÌNH",
            "Năm sinh": 2004,
            "Căn cước": "001204000002",
            "Nghề nghiệp": "Nông dân",
            "Văn hóa": "12/12",
            "Dân tộc": "kinh",
            "Tôn giáo": "không",
            "Nơi ở hiện nay": "Đà Nẵng",
        },
        {
            "STT": 3,
            "TÊN THƯỜNG DÙNG": "PHAN TRẦN GIA HUY",
            "Năm sinh": 2003,
            "Căn cước": "001203000003",
            "Nghề nghiệp": "SINH VIÊN",
            "Văn hóa": "Đại học",
            "Dân tộc": "KINH",
            "Tôn giáo": "Không",
            "Thường trú": "TP Hồ Chí Minh",
        },
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hồ sơ"
    sheet.append(headers)
    sheet.append([])
    for record in records:
        sheet.append([record.get(header) for header in headers])
    workbook.save(path)
    workbook.close()
    return len(headers)


class QuanLyHoSoTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "records.xlsx"
        self.previous_database = app.state["database"]
        app.state["database"] = Path(self.temp_dir.name) / "cache.sqlite3"
        self.header_count = create_test_workbook(self.workbook_path)
        app.set_workbook(self.workbook_path, persist=False)

    def tearDown(self):
        app.state.update(
            workbook=None,
            sheet=None,
            headers={},
            database=self.previous_database,
            record_count=0,
            signature=None,
            loaded_at=None,
            error=None,
            last_output_dir=None,
        )
        self.temp_dir.cleanup()

    def test_reads_records_and_accent_insensitive_search(self):
        self.assertEqual(app.state["record_count"], 3)
        self.assertNotIn("records", app.state)
        matches, count = app.query_records("nguyen", {}, limit=50, offset=0)
        self.assertEqual(count, 1)
        self.assertEqual([record["stt"] for record in matches], ["1"])

        accented_matches, accented_count = app.query_records("Nguyễn", {}, limit=50, offset=0)
        self.assertEqual(accented_count, 1)
        self.assertEqual([record["stt"] for record in accented_matches], ["1"])

    def test_quick_search_does_not_match_unrelated_profile_fields(self):
        matches, count = app.query_records("Ha Noi", {}, limit=50, offset=0)
        self.assertEqual(count, 0)
        self.assertEqual(matches, [])

    def test_filters_ignore_case_and_accents(self):
        occupation = app.query_filter_options("occupation")[0]
        exact_matches, exact_count = app.query_records("", {"occupation": [occupation]}, limit=500, offset=0)
        normalized_matches, normalized_count = app.query_records(
            "", {"occupation": [app.fold(occupation)]}, limit=500, offset=0
        )
        self.assertGreater(exact_count, 0)
        self.assertEqual(exact_count, normalized_count)
        self.assertEqual([record["stt"] for record in exact_matches], [record["stt"] for record in normalized_matches])

    def test_filters_accept_several_values_and_match_any(self):
        _, one = app.query_records("", {"occupation": ["Sinh viên"]}, limit=500, offset=0)
        _, other = app.query_records("", {"occupation": ["Nông dân"]}, limit=500, offset=0)
        matches, both = app.query_records("", {"occupation": ["Sinh viên", "Nông dân"]}, limit=500, offset=0)
        self.assertGreater(other, 0)
        self.assertEqual(both, one + other)
        _, narrowed = app.query_records("", {"occupation": ["Sinh viên", "Nông dân"], "dien": ["__none__"]}, limit=500, offset=0)
        self.assertEqual(narrowed, both)

    def test_index_page_keeps_several_selected_filter_values(self):
        response = app.app.test_client().get("/?occupation=Sinh viên&occupation=Nông dân")
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(html.count('name="occupation" value="Sinh viên" checked'), 1)
        self.assertEqual(html.count('name="occupation" value="Nông dân" checked'), 1)

    def test_list_export_appends_quarter_before_dien(self):
        from quan_ly_ho_so.workbook.list_export import build_list_workbook

        records, _ = app.query_records("", {}, limit=1, offset=0)
        records[0]["quarter"], records[0]["dien"] = "Khu phố 7", "Dân Quân"
        sheet = build_list_workbook(records, app.LIST_TEMPLATE).active
        headings = {cell.value: cell.column for cell in sheet[1] if cell.value}
        self.assertEqual(headings["Diện"], headings["Khu phố"] + 1)
        self.assertEqual(max(headings.values()), headings["Diện"])
        self.assertEqual([sheet.cell(2, headings["Khu phố"]).value, sheet.cell(2, headings["Diện"]).value], ["Khu phố 7", "Dân Quân"])

    def test_related_documents_are_kept_as_files_and_word_becomes_pdf(self):
        import io
        from unittest import mock

        from PIL import Image

        from quan_ly_ho_so import attachments, documents

        picture = io.BytesIO()
        Image.new("RGB", (40, 30), "red").save(picture, "PNG")
        pdf = io.BytesIO()
        Image.new("RGB", (40, 30), "blue").save(pdf, "PDF")

        def fake_word_to_pdf(source, target):
            target.write_bytes(pdf.getvalue())

        def upload(client, csrf, name, content):
            return client.post(
                "/person/1/files",
                data={"csrf_token": csrf, "document": (io.BytesIO(content), name)},
                content_type="multipart/form-data",
            )

        folder = Path(self.temp_dir.name) / "Tài liệu liên quan"
        client = app.app.test_client()
        with client, mock.patch.object(attachments, "_word_to_pdf", fake_word_to_pdf), \
                mock.patch.object(documents, "DOCUMENT_DIR", folder):
            form_text = client.get("/person/1/edit").get_data(as_text=True)
            csrf = re.search(r'name="csrf_token" value="([^"]+)"', form_text).group(1)
            for name, content in (("ảnh.png", picture), ("đơn.pdf", pdf), ("giấy tờ.docx", io.BytesIO(b"word")), ("ảnh.png", picture)):
                self.assertEqual(upload(client, csrf, name, content.getvalue()).status_code, 302)
            owner = folder / "001205000001"
            self.assertEqual(
                {path.name for path in owner.iterdir()}, {"giấy tờ.pdf", "ảnh (2).png", "ảnh.png", "đơn.pdf"}
            )
            self.assertEqual((owner / "ảnh.png").read_bytes(), picture.getvalue())
            self.assertEqual((owner / "đơn.pdf").read_bytes(), pdf.getvalue())
            self.assertEqual((owner / "giấy tờ.pdf").read_bytes(), pdf.getvalue())
            self.assertEqual(client.get("/person/1/documents/ảnh.png").data, picture.getvalue())
            thumbnail = client.get("/person/1/documents/đơn.pdf?size=thumb")
            self.assertEqual(thumbnail.mimetype, "image/jpeg")
            self.assertTrue(thumbnail.data.startswith(b"\xff\xd8"))
            self.assertIn("attachment", client.get("/person/1/documents/ảnh.png?download=1").headers["Content-Disposition"])
            self.assertIn("đơn.pdf", client.get("/person/1/edit").get_data(as_text=True))
            self.assertEqual(client.get("/person/2/documents/ảnh.png").status_code, 404)
            self.assertEqual(client.get("/person/1/documents/..%2F..%2Fsettings.json").status_code, 404)
            client.post("/person/1/documents/ảnh.png/delete", data={"csrf_token": csrf})
            self.assertFalse((owner / "ảnh.png").exists())
            self.assertEqual(upload(client, csrf, "hỏng.png", b"not an image").status_code, 302)
            self.assertFalse((owner / "hỏng.png").exists())
            documents.move_documents("001205000001", "001205000009")
            self.assertFalse(owner.exists())
            self.assertEqual(len(list((folder / "001205000009").iterdir())), 3)

    def test_created_and_updated_dates_track_changes_and_filter(self):
        import sqlite3
        from datetime import date, timedelta

        today = date.today().isoformat()
        yesterday = (date.today() - timedelta(days=1)).isoformat()
        records, count = app.query_records("", {"created_from": today, "created_to": today}, limit=50, offset=0)
        self.assertEqual(count, 3)
        self.assertTrue(all(record["created_at"] and record["updated_at"] for record in records))
        self.assertEqual(app.query_records("", {"created_to": yesterday}, limit=50, offset=0)[1], 0)

        # Pretend everyone was last touched long ago, then change one row from outside the app.
        connection = sqlite3.connect(app.state["database"])
        with connection:
            connection.execute("UPDATE record_meta SET created_at = '2000-01-01T00:00:00', updated_at = '2000-01-01T00:00:00'")
        connection.close()
        workbook = load_workbook(self.workbook_path)
        workbook.active.cell(3, app.state["headers"]["TÊN THƯỜNG DÙNG"]).value = "UPDATED OUTSIDE APP"
        workbook.save(self.workbook_path)
        workbook.close()
        app.refresh_state(force=True)
        recent, recent_count = app.query_records("", {"updated_from": today}, limit=50, offset=0)
        self.assertEqual([record["stt"] for record in recent], ["1"])
        self.assertEqual(app.query_records("", {"created_to": "2000-01-01"}, limit=50, offset=0)[1], 3)

        # Setting a Diện counts as an update; setting the same one again does not.
        app.set_record_type(app.state["workbook"], "001204000002", "DAN_QUAN")
        app.refresh_state(force=True)
        self.assertEqual(app.query_records("", {"updated_from": today}, limit=50, offset=0)[1], 2)

    def test_bulk_delete_hides_records_and_restore_brings_them_back(self):
        client = app.app.test_client()
        with client:
            form_text = client.get("/person/1/edit").get_data(as_text=True)
            csrf = re.search(r'name="csrf_token" value="([^"]+)"', form_text).group(1)
            response = client.post("/bulk/delete", data={"csrf_token": csrf, "stt": ["1", "2"]})
            self.assertEqual(response.status_code, 302)
            self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 1)
            self.assertEqual(app.database_record_count(), 1)
            hidden, hidden_count = app.query_records("", {"deleted": "1"}, limit=50, offset=0)
            self.assertEqual((hidden_count, [record["stt"] for record in hidden]), (2, ["1", "2"]))
            self.assertEqual(app.query_filter_options("occupation"), ["SINH VIÊN"])
            # The workbook still holds the rows, and a resync keeps them hidden.
            app.refresh_state(force=True)
            self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 1)
            self.assertEqual(app.state["record_count"], 1)
            page = client.get("/").get_data(as_text=True)
            self.assertIn('action="/bulk/delete"', page)
            self.assertNotIn('action="/bulk/restore"', page.split("<tbody>")[1])
            page = client.get("/?deleted=1").get_data(as_text=True)
            self.assertIn("Khôi phục", page)
            self.assertIn('action="/bulk/restore"', page.split("<tbody>")[1])
            self.assertNotIn("Xóa hàng loạt", page)
            client.post("/bulk/restore", data={"csrf_token": csrf, "stt": ["1"]})
            self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 2)
            self.assertEqual(client.post("/bulk/delete", data={"csrf_token": "wrong", "stt": ["3"]}).status_code, 302)
            self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 2)

    def test_ca_lists_add_only_new_people_by_citizen_id(self):
        import io

        from quan_ly_ho_so.workbook.ca_import import OTHER_INFO

        ca = Workbook()
        sheet = ca.active
        sheet.append([
            "Số TT", "STT\ntheo diện", "Họ, tên khai sinh\nNgày, tháng, năm sinh", "CCCD", "Nghề nghiệp;\nNơi làm việc",
            "Nơi đăng ký thường trú;\nNơi ở hiện nay của bản thân (tạm trú)", "Dân tộc;\nTôn giáo",
            "Học vấn; Chuyên môn kỹ thuật",
            "Họ tên cha; năm sinh, nghề nghiệp\nHọ tên mẹ; năm sinh, nghề nghiệp\nHọ tên vợ (chồng); năm sinh, nghề nghiệp",
            "Ghi chú", "Khu phố mới", "Khu phố cũ",
        ])
        sheet.append(["3. VẮNG MẶT ĐỊA PHƯƠNG"])
        sheet.append([1, 1, "Lê Văn Mới\n02/11/2002", "079202000111", "Thợ hàn\nCông ty A", "1 Phố A, KP 1\n2 Phố B", "Kinh\nKhông",
                      "12/12\nTrung cấp\nHàn", "Lê Văn Cha 1970,\nNông dân\nTrần Thị Mẹ 1972,\nBuôn bán", "Sức khỏe tốt", 7, 27])
        sheet.append([2, 2, "Trần Đã Có\n01/01/2001", "001205000001", None, None, "Kinh\nKhông", None, None, None, 8, 28])
        sheet.append([3, 3, "Không Có Căn Cước\n05/05/2003", None, None, "3 Phố C", None, None, None, None, 9, 29])
        sheet.append([4, 4, "Phạm Trùng Lặp\n03/03/2000", "079200000222", "Sinh viên", None, None, "12/12", "Phạm Một Mình 1960,\nHưu trí", None, 10, 30])
        sheet.append([5, 5, "Phạm Trùng Lặp\n03/03/2000", "079200000222", None, None, None, None, None, None, 10, 30])
        content = io.BytesIO()
        ca.save(content)

        client = app.app.test_client()
        with client:
            form_text = client.get("/person/1/edit").get_data(as_text=True)
            csrf = re.search(r'name="csrf_token" value="([^"]+)"', form_text).group(1)
            response = client.post(
                "/import/ca",
                data={"csrf_token": csrf, "ca_files": (io.BytesIO(content.getvalue()), "kp7.xlsx")},
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 302)
            # Step one only reads the file: nothing is written until the manager confirms.
            preview = client.get("/").get_data(as_text=True)
            self.assertIn("chưa thêm gì vào Excel tổng", preview)
            self.assertIn("2 sẽ được thêm mới", preview)
            self.assertIn("Thêm 2 người vào Excel tổng", preview)
            self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 3)
            self.assertEqual(list(self.workbook_path.parent.glob("records.backup-*.xlsx")), [])
            client.post("/import/ca/cancel", data={"csrf_token": csrf})
            self.assertNotIn("chưa thêm gì", client.get("/").get_data(as_text=True))
            self.assertEqual(client.post("/import/ca/confirm", data={"csrf_token": csrf}).status_code, 302)
            self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 3)

            client.post(
                "/import/ca",
                data={"csrf_token": csrf, "ca_files": (io.BytesIO(content.getvalue()), "kp7.xlsx")},
                content_type="multipart/form-data",
            )
            self.assertEqual(client.post("/import/ca/confirm", data={"csrf_token": csrf}).status_code, 302)
            page = client.get("/").get_data(as_text=True)
            self.assertIn("2 thêm mới", page)
            self.assertIn("1 đã có", page)
            self.assertIn("Không Có Căn Cước", page)
            self.assertNotIn("Kết quả nhập file CA", client.get("/").get_data(as_text=True))

        self.assertEqual(app.query_records("", {}, limit=50, offset=0)[1], 5)
        added = app.query_records("", {"created_from": "2000-01-01"}, limit=50, offset=0)[0][-2:]
        first, second = added
        self.assertEqual((first["stt"], first["name"], first["citizen_id"], first["quarter"]), ("4", "LÊ VĂN MỚI", "079202000111", "Khu phố 7"))
        self.assertEqual(first["data"]["Tên khai sinh"], "Lê Văn Mới")
        self.assertEqual((first["data"]["Ngày sinh"], first["data"]["Tháng sinh"], first["data"]["Năm sinh"]), ("2", "11", "2002"))
        self.assertEqual((first["data"]["Nghề nghiệp"], first["data"]["Nơi làm việc, học tập"]), ("Thợ hàn", "Công ty A"))
        self.assertEqual((first["data"]["Thường trú"], first["data"]["Nơi ở hiện nay"]), ("1 Phố A, KP 1", "2 Phố B"))
        self.assertEqual((first["data"]["Văn hóa"], first["data"]["Chuyên môn"], first["data"]["Ngành đào tạo"]), ("12/12", "Trung cấp", "Hàn"))
        self.assertEqual((first["data"]["Tên cha"], first["data"]["năm sinh cha"], first["data"]["Nghề nghiệp cha"]), ("Lê Văn Cha", "1970", "Nông dân"))
        self.assertEqual((first["data"]["Tên mẹ"], first["data"]["Năm sinh mẹ"]), ("Trần Thị Mẹ", "1972"))
        # A single parent line cannot be told apart, so it stays as written instead of being guessed.
        self.assertEqual(second["data"].get("Tên cha", ""), "")
        self.assertIn("Phạm Một Mình 1960,", second["data"].get(OTHER_INFO, ""))
        self.assertEqual(app.find_record("2")["name"], "TRẦN THỊ BÌNH")  # existing person untouched
        self.assertEqual(len(list(self.workbook_path.parent.glob("records.backup-*.xlsx"))), 1)

    def test_filter_options_group_case_variations(self):
        options = app.query_filter_options("ethnicity")
        kinh_options = [value for value in options if app.fold(value) == "kinh"]
        self.assertEqual(len(kinh_options), 1)

    def test_external_excel_change_is_synchronized(self):
        workbook = load_workbook(self.workbook_path)
        name_column = app.state["headers"]["TÊN THƯỜNG DÙNG"]
        workbook.active.cell(3, name_column).value = "UPDATED OUTSIDE APP"
        workbook.save(self.workbook_path)
        workbook.close()

        self.assertTrue(app.refresh_state())
        self.assertEqual(app.find_record("1")["name"], "UPDATED OUTSIDE APP")

    def test_failed_sync_keeps_previous_sqlite_data(self):
        self.workbook_path.write_bytes(b"not an Excel workbook")
        self.assertFalse(app.refresh_state(force=True))
        self.assertEqual(app.database_record_count(), 3)
        self.assertIsNotNone(app.find_record("1"))

    def test_edit_and_add_preserve_workbook_shape(self):
        signature = app.signature_token(app.state["signature"])
        name_column = app.state["headers"]["TÊN THƯỜNG DÙNG"]
        row, stt, backup = app.save_person(
            self.workbook_path, app.state["sheet"], {name_column: "EDITED TEST"}, signature, original_stt="1"
        )
        self.assertEqual((row, stt), (3, "1"))
        self.assertTrue(backup.is_file())

        app.set_workbook(self.workbook_path, persist=False)
        signature = app.signature_token(app.state["signature"])
        row, stt, backup = app.save_person(
            self.workbook_path, app.state["sheet"], {name_column: "ADDED TEST"}, signature
        )
        self.assertEqual((row, stt), (6, "4"))
        self.assertTrue(backup.is_file())
        workbook = load_workbook(self.workbook_path, read_only=True, data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet.cell(3, name_column).value, "EDITED TEST")
        self.assertEqual(sheet.cell(6, name_column).value, "ADDED TEST")
        self.assertEqual(sheet.cell(6, app.state["headers"]["STT"]).value, 4)
        self.assertEqual(sheet.max_column, self.header_count)
        workbook.close()

    def test_stale_form_is_rejected(self):
        signature = app.signature_token(app.state["signature"])
        workbook = load_workbook(self.workbook_path)
        workbook.active["C3"] = "CHANGED OUTSIDE APP"
        workbook.save(self.workbook_path)
        workbook.close()
        with self.assertRaises(app.StaleWorkbookError):
            app.save_person(
                self.workbook_path,
                app.state["sheet"],
                {app.state["headers"]["TÊN THƯỜNG DÙNG"]: "STALE EDIT"},
                signature,
                original_stt="1",
            )

    def test_browser_routes_and_word_generation(self):
        client = app.app.test_client()
        with client:
            response = client.get("/")
            self.assertEqual(response.status_code, 200)
            self.assertIn(b"PHAN TR", response.data)
            page = response.get_data(as_text=True)
            self.assertIn("Tìm kiếm hồ sơ", page)
            self.assertIn("Bộ lọc nâng cao", page)
            self.assertIn("Thêm hồ sơ", page)
            response = client.get("/?q=nguyen&birth_year=2005")
            self.assertEqual(response.status_code, 200)
            self.assertIn('<details class="advanced" open>', response.get_data(as_text=True))
            form = client.get("/person/1/edit")
            self.assertEqual(form.status_code, 200)
            form_text = form.get_data(as_text=True)
            self.assertIn("Cập nhật hồ sơ", form_text)
            self.assertIn("Thông tin cá nhân", form_text)
            self.assertIn("Lưu vào Excel", form_text)
            self.assertIn('<details class="section-card" open>', form_text)
            csrf = re.search(r'name="csrf_token" value="([^"]+)"', form_text).group(1)
            response = client.post("/generate/1", data={"csrf_token": csrf}, follow_redirects=False)
            self.assertEqual(response.status_code, 302)
        output_dir = self.workbook_path.parent / app.DEFAULT_OUTPUT_NAME
        first = generate_document(self.workbook_path, app.state["sheet"], 3, TEMPLATE, output_dir)
        second = generate_document(self.workbook_path, app.state["sheet"], 3, TEMPLATE, output_dir)
        self.assertTrue(first.is_file())
        self.assertTrue(second.is_file())
        self.assertNotEqual(first, second)


if __name__ == "__main__":
    unittest.main()
